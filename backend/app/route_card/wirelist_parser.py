"""Parse BEL Wire List (WL) PDF text into structured pin/wire rows."""

from __future__ import annotations

import re

from app.route_card.pdf_ingest import ingest_pdf, tables_to_tsv_lines

WIRE_ROW = re.compile(
    r"(?P<sl>\d+)\s+"
    r"(?P<frm>(?:J\d+-\d+|P\d+-[A-Z]))\s+"
    r"(?P<to>PK\d+-P\d+-\d+)\s+"
    r"(?P<fn>[A-Z0-9_]+)\s+"
    r"(?P<wtype>WIRE\s+LFH[^\d]*?)\s+"
    r"(?P<pno>\d{4}\s+\d{3}\s+\d{3}\s+\d{2})\s+"
    r"(?P<length>\d{3,5})"
    r"(?:\s+(?P<remarks>.+))?",
    re.I,
)
TITLE_PART = re.compile(
    r"(?:PART\s*(?:NUMBER|NO\.?|NUM)|P/?N)\s*\n\s*(\d{4}\s+\d{3}\s+\d{3}\s+\d{2})",
    re.I,
)
TITLE_ALT = re.compile(
    r"(\d{4}\s+\d{3}\s+\d{3}\s+\d{2})\s*\n(?:PART\s*(?:NUMBER|NO\.?|NUM)|P/?N)",
    re.I,
)
DOC_TITLE = re.compile(
    r"(?:TITLE|PART\s*NAME|DESIGNATION)\s*\n(.+?)(?:\n\d{4}\s+\d{3}|\nPART\s*(?:NUMBER|NO)|P/?N|\nHELD)",
    re.S | re.I,
)
NOTE_PL = re.compile(
    r"(?:PL|PARTS?\s*LIST)\s*item\s*(?:no\.?|number|#)?\s*:?\s*0*(\d+)",
    re.I,
)
PART_NO_RE = re.compile(r"(\d{4}\s+\d{3}\s+\d{3}\s+\d{2})")
PART_NO_FLEX = re.compile(r"(\d{4}\s+\d{3}\s+\d{1,3}(?:\s+\d{1,2})?)")
FROM_RE = re.compile(r"^(?:J\d+-\d+|P\d+-[A-Z0-9]+)$", re.I)
TO_RE = re.compile(r"^PK\d+-P\d+-\s*\d+$", re.I)
WL_NOTE_LINE = re.compile(r"(?:notes?|remarks?)\s*\d*\s*:?", re.I)


def parse_wirelist_text(text: str, filename: str = "", *, format_aliases: dict | None = None) -> dict:
    if format_aliases:
        from app.route_card.format_templates import apply_aliases_to_text

        text = apply_aliases_to_text(text, format_aliases, doc_type="WL")
    warnings: list[str] = []
    part_number = ""
    m = TITLE_PART.search(text) or TITLE_ALT.search(text)
    if m:
        part_number = re.sub(r"\s+", " ", m.group(1)).strip()
    elif filename:
        digits = re.sub(r"\D", "", filename)
        if len(digits) >= 12:
            part_number = f"{digits[0:4]} {digits[4:7]} {digits[7:10]} {digits[10:12]}"

    title = "Wire List"
    tm = DOC_TITLE.search(text)
    if tm:
        title = re.sub(r"\s+", " ", tm.group(1)).strip()[:120]
    if "CABLE ASSY" in text.upper():
        cm = re.search(r"CABLE ASSY[^\n]{0,80}", text, re.I)
        if cm:
            title = re.sub(r"\s+", " ", cm.group(0)).strip()

    flat = re.sub(r"[ \t]+", " ", text)
    joined = re.sub(r"\n+", "\n", flat)

    wires = []
    for m in WIRE_ROW.finditer(joined.replace("\n", " ")):
        wires.append(
            {
                "slNo": int(m.group("sl")),
                "from": m.group("frm"),
                "to": m.group("to"),
                "function": m.group("fn"),
                "wireType": re.sub(r"\s+", " ", m.group("wtype")).strip(),
                "partNo": re.sub(r"\s+", " ", m.group("pno")).strip(),
                "lengthMm": int(m.group("length")),
                "remarks": (m.group("remarks") or "").strip()[:200],
            }
        )

    if len(wires) < 10:
        wires = _loose_rows(text)
        if not wires:
            warnings.append("Wire List text found but few/no wire rows parsed. Prefer a clean PDF export.")

    materials: dict[str, dict] = {}
    for w in wires:
        key = w["partNo"]
        mat = materials.setdefault(
            key,
            {"partNo": key, "wireType": w["wireType"], "count": 0, "totalLengthMm": 0},
        )
        mat["count"] += 1
        mat["totalLengthMm"] += int(w.get("lengthMm") or 0)

    notes = []
    for line in text.splitlines():
        low = line.lower()
        if (
            WL_NOTE_LINE.search(line)
            or "fold back" in low
            or "twist" in low
            or "note 1" in low
            or "remark" in low
        ):
            cleaned = re.sub(r"\s+", " ", line).strip()
            if len(cleaned) >= 8:
                notes.append(cleaned)
    pl_refs = [{"itemNo": int(x), "source": "WL note"} for x in NOTE_PL.findall(text)]

    return {
        "docType": "WL",
        "partNumber": part_number,
        "title": title,
        "wireCount": len(wires),
        "wires": wires,
        "materials": list(materials.values()),
        "notes": notes,
        "plItemRefs": pl_refs,
        "warnings": warnings,
        "raw_text": text[:100000],
    }


def _loose_rows(text: str) -> list[dict]:
    rows = []
    pat = re.compile(
        r"(\d+)\s+(J\d+-\d+|P\d+-[A-Z])\s+(PK\d+-P\d+-\s*\d+)\s+(\S+)\s+.*?(\d{4}\s+\d{3}\s+\d{3}\s+\d{2})\s+(\d{3,5})",
        re.I,
    )
    for m in pat.finditer(re.sub(r"\s+", " ", text)):
        rows.append(
            {
                "slNo": int(m.group(1)),
                "from": m.group(2),
                "to": re.sub(r"\s+", "", m.group(3)),
                "function": m.group(4),
                "wireType": "WIRE LFH",
                "partNo": re.sub(r"\s+", " ", m.group(5)).strip(),
                "lengthMm": int(m.group(6)),
                "remarks": "",
            }
        )
    return rows


def _wires_from_tables(tables: list[dict]) -> list[dict]:
    """Map pdfplumber table rows that look like wire-list lines."""
    rows: list[dict] = []
    for t in tables:
        for raw in t.get("rows") or []:
            cells = [re.sub(r"\s+", " ", str(c or "")).strip() for c in raw if str(c or "").strip()]
            if len(cells) < 5:
                continue
            if not re.fullmatch(r"\d+", cells[0]):
                continue
            frm = next((c for c in cells[1:] if FROM_RE.match(c)), "")
            to = next((c for c in cells[1:] if TO_RE.match(re.sub(r"\s+", "", c)) or TO_RE.match(c)), "")
            pn = next((c for c in cells if PART_NO_RE.search(c) or PART_NO_FLEX.search(c)), "")
            if not (frm and to and pn):
                continue
            pn_m = PART_NO_RE.search(pn) or PART_NO_FLEX.search(pn)
            length = None
            for c in reversed(cells):
                if re.fullmatch(r"\d{3,5}", c):
                    length = int(c)
                    break
            if length is None:
                continue
            fn = ""
            for c in cells[1:]:
                if c in {frm, to, pn} or PART_NO_RE.search(c) or FROM_RE.match(c) or TO_RE.match(c):
                    continue
                if re.fullmatch(r"\d{3,5}", c):
                    continue
                if "WIRE" in c.upper():
                    continue
                if re.fullmatch(r"[A-Z0-9_]+", c, re.I) and not fn:
                    fn = c
            wtype = next((c for c in cells if "WIRE" in c.upper()), "WIRE LFH")
            remarks = ""
            for c in cells:
                if c.lower().startswith("twist") or "note" in c.lower():
                    remarks = c
                    break
            rows.append(
                {
                    "slNo": int(cells[0]),
                    "from": frm,
                    "to": re.sub(r"\s+", "", to),
                    "function": fn,
                    "wireType": re.sub(r"\s+", " ", wtype).strip(),
                    "partNo": re.sub(r"\s+", " ", pn_m.group(1)).strip() if pn_m else pn,
                    "lengthMm": length,
                    "remarks": remarks[:200],
                }
            )
    return rows


def _excel_to_tsv_text(data: bytes, ext: str) -> tuple[str, list[str]]:
    """Flatten .xls/.xlsx sheets to tab-separated text for regex WL parsing."""
    warnings: list[str] = []
    lines: list[str] = []
    if ext == "xls":
        head = data[:64]
        if not head.startswith(b"\xd0\xcf") and (b"\t" in head or b"," in head):
            try:
                return data.decode("utf-8", errors="ignore"), ["Wire List .xls was plain text/TSV."]
            except Exception:
                return data.decode("latin-1", errors="ignore"), ["Wire List .xls was plain text/TSV."]
        try:
            import xlrd
        except ImportError:
            return "", ["xlrd is not installed; cannot parse Wire List .xls."]
        try:
            book = xlrd.open_workbook(file_contents=data)
        except Exception as exc:
            try:
                return data.decode("utf-8", errors="ignore"), [f"Could not open Wire List as Excel ({exc})."]
            except Exception:
                return "", [f"Could not open Wire List .xls: {exc}"]
        for si in range(book.nsheets):
            sheet = book.sheet_by_index(si)
            for r in range(sheet.nrows):
                cells = [str(sheet.cell_value(r, c)).strip() for c in range(sheet.ncols)]
                if any(cells):
                    lines.append("\t".join(cells))
        return "\n".join(lines), warnings

    if ext == "xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError:
            return "", ["openpyxl is not installed; cannot parse Wire List .xlsx."]
        from io import BytesIO

        try:
            book = load_workbook(BytesIO(data), read_only=True, data_only=True)
        except Exception as exc:
            return "", [f"Could not open Wire List .xlsx: {exc}"]
        try:
            for sheet in book.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    cells = [str(c or "").strip() for c in (row or [])]
                    if any(cells):
                        lines.append("\t".join(cells))
        finally:
            book.close()
        return "\n".join(lines), warnings

    return "", [f"Unsupported Wire List type .{ext}."]


def parse_wirelist_bytes(data: bytes, filename: str = "", *, format_aliases: dict | None = None) -> dict:
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext in {"xls", "xlsx"}:
        text, warnings = _excel_to_tsv_text(data, ext)
        if not (text or "").strip():
            return {
                "docType": "WL",
                "partNumber": "",
                "title": filename,
                "wireCount": 0,
                "wires": [],
                "materials": [],
                "notes": [],
                "plItemRefs": [],
                "warnings": warnings or ["Wire List Excel file could not be read."],
                "raw_text": "",
                "ingest_method": "excel",
            }
        parsed = parse_wirelist_text(text, filename, format_aliases=format_aliases)
        parsed["warnings"] = list(dict.fromkeys(warnings + (parsed.get("warnings") or [])))
        parsed["ingest_method"] = "excel"
        return parsed

    if ext != "pdf":
        return {
            "docType": "WL",
            "partNumber": "",
            "title": filename,
            "wireCount": 0,
            "wires": [],
            "materials": [],
            "notes": [],
            "plItemRefs": [],
            "warnings": [
                f"Unsupported Wire List type .{ext or 'file'}. Upload a PDF or Excel file."
            ],
            "raw_text": "",
            "ingest_method": "none",
        }

    result = ingest_pdf(data, want_tables=True)
    warnings = list(result.warnings)
    if len(result.text.strip()) < 40 and not result.tables:
        return {
            "docType": "WL",
            "partNumber": "",
            "title": filename,
            "wireCount": 0,
            "wires": [],
            "materials": [],
            "notes": [],
            "plItemRefs": [],
            "warnings": warnings
            + (["Wire List PDF has little/no text layer."] if not warnings else []),
            "raw_text": result.text[:5000],
            "ingest_method": result.method,
        }

    # Enrich text with table TSV so existing regex can still match.
    text = result.text
    tsv = tables_to_tsv_lines(result.tables)
    if tsv:
        text = f"{text}\n\n{tsv}"

    parsed = parse_wirelist_text(text, filename, format_aliases=format_aliases)
    table_wires = _wires_from_tables(result.tables)
    if table_wires and len(table_wires) > len(parsed.get("wires") or []):
        # Prefer denser table parse when regex under-matched.
        materials: dict[str, dict] = {}
        for w in table_wires:
            key = w["partNo"]
            mat = materials.setdefault(
                key,
                {"partNo": key, "wireType": w["wireType"], "count": 0, "totalLengthMm": 0},
            )
            mat["count"] += 1
            mat["totalLengthMm"] += int(w.get("lengthMm") or 0)
        parsed["wires"] = table_wires
        parsed["wireCount"] = len(table_wires)
        parsed["materials"] = list(materials.values())
        parsed["warnings"] = [
            w for w in (parsed.get("warnings") or []) if "few/no wire rows" not in w.lower()
        ]

    parsed["warnings"] = list(dict.fromkeys(warnings + (parsed.get("warnings") or [])))
    if result.method == "ocr":
        parsed["warnings"].insert(0, "Wire List used offline OCR (Tesseract).")
    parsed["ingest_method"] = result.method
    parsed["tableCount"] = len(result.tables)
    return parsed


def wirelist_to_route_additions(wl: dict, start_op: int = 1000) -> dict:
    """Extra operations + inspection + materials from a parsed WL."""
    ops = []
    op_no = start_op
    wires = wl.get("wires") or []
    materials = wl.get("materials") or []

    if materials:
        mat_lines = "; ".join(
            f"{m['partNo']} x{m['count']} ({m['totalLengthMm']} mm)" for m in materials[:12]
        )
        ops.append(
            {
                "id": f"op-{op_no}",
                "opNo": op_no,
                "operation": "Wire cutting / preparation",
                "workCentre": "Cable Prep",
                "machine": "Cable bench",
                "tool": "Wire cutter / stripper",
                "items": mat_lines[:200],
                "itemsUsed": [
                    {
                        "itemNo": None,
                        "partNo": m["partNo"],
                        "qty": m["count"],
                        "description": m.get("wireType"),
                    }
                    for m in materials
                ],
                "torqueSpec": "",
                "reference": wl.get("partNumber") or "WL",
                "setup": "1",
                "time": "",
                "inspection": "Length check vs WL",
                "instructionText": f"Cut and prepare wires per WL {wl.get('partNumber') or ''}: {mat_lines}",
                "status": "Planned",
                "type": "Cable Prep",
                "reason": "From Wire List materials",
                "features": [],
            }
        )
        op_no += 10

    twist = sum(1 for w in wires if "twist" in (w.get("remarks") or "").lower())
    if twist:
        ops.append(
            {
                "id": f"op-{op_no}",
                "opNo": op_no,
                "operation": "Twisted-pair forming",
                "workCentre": "Cable Prep",
                "machine": "Cable bench",
                "tool": "Hand tools",
                "items": "",
                "itemsUsed": [],
                "torqueSpec": "",
                "reference": "WL remarks",
                "setup": "1",
                "time": "",
                "inspection": "Visual — pairs twisted",
                "instructionText": f"Form twisted pairs for {twist} wire(s) marked in the Wire List.",
                "status": "Planned",
                "type": "Cable Prep",
                "reason": "WL twisted-pair remarks",
                "features": [],
            }
        )
        op_no += 10

    if wires:
        ops.append(
            {
                "id": f"op-{op_no}",
                "opNo": op_no,
                "operation": "Connector termination / pin map",
                "workCentre": "Assembly",
                "machine": "Assembly station",
                "tool": "Crimp / solder tools",
                "items": f"{len(wires)} terminations",
                "itemsUsed": [],
                "torqueSpec": "",
                "reference": wl.get("title") or "WL",
                "setup": "1",
                "time": "",
                "inspection": "Continuity 100%",
                "instructionText": (
                    f"Terminate {len(wires)} wires per Wire List pin map "
                    f"({wl.get('title') or wl.get('partNumber') or 'WL'})."
                ),
                "status": "Planned",
                "type": "Assembly",
                "reason": "WL pin map",
                "features": [],
            }
        )
        op_no += 10

    for note in wl.get("notes") or []:
        ops.append(
            {
                "id": f"op-{op_no}",
                "opNo": op_no,
                "operation": "WL special instruction",
                "workCentre": "Assembly",
                "machine": "Assembly station",
                "tool": "—",
                "items": "",
                "itemsUsed": [],
                "torqueSpec": "",
                "reference": "WL Note",
                "setup": "1",
                "time": "",
                "inspection": "Visual",
                "instructionText": note,
                "status": "Planned",
                "type": "Assembly",
                "reason": note[:120],
                "features": [],
            }
        )
        op_no += 10

    inspection = []
    for w in wires[:40]:
        inspection.append(
            {
                "id": f"wl-ic-{w['slNo']}",
                "dimension": f"Continuity {w['from']} → {w['to']}",
                "nominal": w.get("function") or "Pass",
                "tolerance": "Pass/Fail",
                "method": "Continuity tester",
                "frequency": "100%",
                "criticality": "Important",
            }
        )
    if len(wires) > 40:
        inspection.append(
            {
                "id": "wl-ic-rest",
                "dimension": f"Continuity remaining {len(wires) - 40} nets",
                "nominal": "Pass",
                "tolerance": "Pass/Fail",
                "method": "Continuity tester",
                "frequency": "100%",
                "criticality": "Important",
            }
        )

    return {"operations": ops, "inspection": inspection, "materials": materials}
